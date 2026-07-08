import streamlit as st
import docx
from io import BytesIO

# Configurazione pagina
st.set_page_config(page_title="Motore LLM - Pratiche Edilizie", layout="wide")
st.title("Generatore Documenti Pratiche Edilizie")

# ==========================================
# GESTIONE STATO (SESSION STATE)
# ==========================================
# Inizializzazione liste di ID univoci per gestire aggiunte e rimozioni
if 'intestatari_ids' not in st.session_state:
    st.session_state.intestatari_ids = [0] # Il primo intestatario (ID 0) c'è sempre
if 'professionisti_ids' not in st.session_state:
    st.session_state.professionisti_ids = [] # Lista vuota in partenza per gli extra
if 'counter' not in st.session_state:
    st.session_state.counter = 1 # Contatore per generare ID sempre nuovi

# Funzioni di utilità per aggiungere/rimuovere
def add_intestatario():
    st.session_state.intestatari_ids.append(st.session_state.counter)
    st.session_state.counter += 1

def remove_intestatario(id_to_remove):
    st.session_state.intestatari_ids.remove(id_to_remove)

def add_professionista():
    st.session_state.professionisti_ids.append(st.session_state.counter)
    st.session_state.counter += 1

def remove_professionista(id_to_remove):
    st.session_state.professionisti_ids.remove(id_to_remove)


# ==========================================
# SEZIONE 1: RACCOGLITORE DI DATI
# ==========================================
st.header("1. Anagrafiche e Raccoglitore Dati")

# --- INTESTATARI ---
st.subheader("Intestatari")
intestatari_data = []

for idx, uid in enumerate(st.session_state.intestatari_ids):
    st.markdown(f"**Intestatario {idx+1}**")
    
    col1, col2, col3, col4 = st.columns(4)
    with col1:
        nome = st.text_input("Nome", key=f"int_nome_{uid}")
        cognome = st.text_input("Cognome", key=f"int_cogn_{uid}")
        cf = st.text_input("Codice Fiscale", key=f"int_cf_{uid}")
    with col2:
        data_nascita = st.date_input("Data di Nascita", key=f"int_data_{uid}")
        luogo_nascita = st.text_input("Luogo di Nascita", key=f"int_luogo_{uid}")
        titolo = st.selectbox("Diritto", ["Proprietario", "Comproprietario", "Usufruttuario", "Rappresentante Legale"], key=f"int_tit_{uid}")
    with col3:
        via = st.text_input("Indirizzo di residenza (Via)", key=f"int_via_{uid}")
        civico = st.text_input("Civico", key=f"int_civ_{uid}")
        cap = st.text_input("CAP", key=f"int_cap_{uid}")
    with col4:
        paese = st.text_input("Paese", key=f"int_paese_{uid}")
        prov = st.text_input("Provincia", key=f"int_prov_{uid}")
        mail = st.text_input("Mail", key=f"int_mail_{uid}")
        tel = st.text_input("Telefono", key=f"int_tel_{uid}")
    
    # Bottone di rimozione (solo se non è il primo intestatario)
    if idx > 0:
        st.button("🗑️ Rimuovi questo intestatario", key=f"del_int_{uid}", on_click=remove_intestatario, args=(uid,))
    
    # Salvataggio dati per il Word
    intestatari_data.append({
        "Nome": nome, 
        "Cognome": cognome, 
        "CF": cf, 
        "Data Nascita": data_nascita.strftime('%d/%m/%Y') if data_nascita else "", 
        "Luogo Nascita": luogo_nascita,
        "Residenza": f"{via} {civico}, {cap} {paese} ({prov})",
        "Mail": mail, 
        "Tel": tel, 
        "Diritto": titolo
    })
    st.divider()

st.button("➕ Aggiungi un altro intestatario", on_click=add_intestatario)

# --- PROFESSIONISTI ---
st.subheader("Professionisti")

# Funzione per renderizzare il form del professionista
def render_professionista(uid, is_main=False):
    if is_main:
        qualifica = "Progettista"
        st.markdown(f"**Professionista 1 - {qualifica} (Principale)**")
    else:
        qualifica = st.selectbox("Qualifica", ["Strutturista", "Termotecnico", "DDLL", "Coordinatore Sicurezza", "Collaudatore"], key=f"prof_qual_{uid}")
    
    col1, col2, col3 = st.columns(3)
    with col1:
        tit_opzioni = ["Architetto", "Ingegnere", "Geometra", "Perito", "Altro"]
        titolo = st.selectbox("Titolo", tit_opzioni, index=0, key=f"prof_tit_{uid}")
        nome = st.text_input("Nome", value="Gianfranco" if is_main else "", key=f"prof_nome_{uid}")
        cognome = st.text_input("Cognome", value="Andriolo" if is_main else "", key=f"prof_cogn_{uid}")
        cf = st.text_input("Codice Fiscale", value="NDRGFR83S17F964E" if is_main else "", key=f"prof_cf_{uid}")
    with col2:
        luogo_nascita = st.text_input("Luogo di Nascita", value="Noventa Vicentina" if is_main else "", key=f"prof_luogo_{uid}")
        data_nascita = st.text_input("Data di Nascita (gg/mm/aaaa)", value="17/11/1983" if is_main else "", key=f"prof_data_{uid}")
        studio = st.text_input("Studio in (Indirizzo)", value="Via Masotto 11, 36025 Noventa Vicentina" if is_main else "", key=f"prof_studio_{uid}")
    with col3:
        pec = st.text_input("PEC", value="andriolo.18135@oamilano.it" if is_main else "", key=f"prof_pec_{uid}")
        mail = st.text_input("Mail", value="studioandriolo@gmail.com" if is_main else "", key=f"prof_mail_{uid}")
        tel = st.text_input("Telefono", value="377 9662445" if is_main else "", key=f"prof_tel_{uid}")
        albo_coll = st.text_input("Iscrizione all'Albo/Collegio di", value="Milano" if is_main else "", key=f"prof_albo_{uid}")
        num_albo = st.text_input("N. di Iscrizione", value="18135" if is_main else "", key=f"prof_num_{uid}")
        
    if not is_main:
        st.button("🗑️ Rimuovi professionista", key=f"del_prof_{uid}", on_click=remove_professionista, args=(uid,))
    
    st.divider()
    
    # Ritorna il dizionario completo per il Word
    return {
        "Qualifica": qualifica, 
        "Nome": f"{titolo} {nome} {cognome}", 
        "CF": cf,
        "Dati": f"Nato a {luogo_nascita} il {data_nascita} | Studio: {studio} | PEC: {pec} | Mail: {mail} | Tel: {tel} | Albo: {albo_coll} n. {num_albo}"
    }

professionisti_data = []

# Disegna Professionista Principale (ID fisso 'main')
prof_main = render_professionista('main', is_main=True)
professionisti_data.append(prof_main)

# Disegna Professionisti Extra
for uid in st.session_state.professionisti_ids:
    prof_extra = render_professionista(uid, is_main=False)
    professionisti_data.append(prof_extra)

st.button("➕ Aggiungi altro professionista", on_click=add_professionista)


# ==========================================
# SEZIONE 2: ELEMENTI COGNITIVI
# ==========================================
st.header("2. Elementi Cognitivi del Progetto")
tipo_pratica = st.selectbox("Tipo di pratica edilizia", ["CILA", "SCIA art.22", "SCIA art.23", "PdC"])
tipo_intervento = st.text_input("Tipo di intervento (es. Manutenzione Straordinaria, Ristrutturazione)")
rel_sintetica = st.text_area("Relazione Sintetica dell'intervento", height=150)
elaborati = st.file_uploader("Elaborati grafici (Carica PDF)", type=['pdf'], accept_multiple_files=True)


# ==========================================
# SEZIONE 3: GENERAZIONE DOCUMENTI (LLM)
# ==========================================
st.header("3. Documenti da Redigere")
st.caption("Seleziona i documenti necessari. Per ognuno, carica un documento di spunto e il template .docx da compilare.")

docs_necessari = {
    "Relazione tecnico-illustrativa": st.checkbox("Relazione tecnico-illustrativa", value=True),
    "Relazione paesaggistica": st.checkbox("Relazione paesaggistica"),
    "Relazione stato legittimo": st.checkbox("Relazione stato legittimo"),
    "Modello FCA (P/P/P/I/A)": st.checkbox("Modello FCA (P/P/P/I/A)"),
    "Relazione FLDm": st.checkbox("Relazione fattore luce diurna media (FLDm)"),
    "Relazione calcolo volumi": st.checkbox("Relazione sul calcolo dei volumi")
}

file_templates = {}

for doc_name, is_checked in docs_necessari.items():
    if is_checked:
        with st.expander(f"Configura: {doc_name}", expanded=True):
            col1, col2 = st.columns(2)
            with col1:
                spunto = st.file_uploader(f"Spunto per {doc_name} (PDF/DOCX)", key=f"spunto_{doc_name}")
            with col2:
                template = st.file_uploader(f"Template da compilare per {doc_name} (.docx)", type=['docx'], key=f"tpl_{doc_name}")
            
            file_templates[doc_name] = {"spunto": spunto, "template": template}


# ==========================================
# LOGICA DI ESPORTAZIONE
# ==========================================
st.divider()

def genera_docx_anagrafica():
    doc = docx.Document()
    doc.add_heading('Dati per Portale "Impresa in un giorno"', 0)
    
    # Sezione Intestatari
    doc.add_heading('Intestatari', level=1)
    for idx, int_data in enumerate(intestatari_data):
        doc.add_heading(f'Soggetto {idx+1} - {int_data["Diritto"]}', level=2)
        for key, value in int_data.items():
            doc.add_paragraph(f"{key}: {value}", style='List Bullet')
            
    # Sezione Professionisti
    doc.add_heading('Professionisti', level=1)
    for prof in professionisti_data:
        doc.add_heading(f'{prof["Qualifica"]}: {prof["Nome"]}', level=2)
        doc.add_paragraph(f"C.F.: {prof['CF']}", style='List Bullet')
        doc.add_paragraph(f"Dati aggiuntivi: {prof['Dati']}", style='List Bullet')
        
    # Sezione Progetto
    doc.add_heading('Dati Progetto', level=1)
    doc.add_paragraph(f"Pratica: {tipo_pratica}", style='List Bullet')
    doc.add_paragraph(f"Intervento: {tipo_intervento}", style='List Bullet')
    doc.add_paragraph(f"Sintesi: {rel_sintetica}", style='List Bullet')
    
    # Salva in memoria
    io_stream = BytesIO()
    doc.save(io_stream)
    return io_stream.getvalue()

col_btn1, col_btn2 = st.columns(2)

with col_btn1:
    st.download_button(
        label="📄 Scarica File Dati (Copia-Incolla)",
        data=genera_docx_anagrafica(),
        file_name="dati_portale_impresainungiorno.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )

with col_btn2:
    if st.button("🤖 Genera Relazioni con LLM"):
        st.info("Logica LLM/Mail Merge da implementare in base alla strada che sceglieremo.")
